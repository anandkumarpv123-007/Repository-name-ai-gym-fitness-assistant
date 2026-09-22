import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout_box(doc, text, title="IMPORTANT SAFETY NOTICE", bg_hex="FFFBEB", border_hex="F59E0B"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set left border thick
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')  # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_hex)
    tcBorders.append(left)

    for side in ['top', 'bottom', 'right']:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)

    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"🛡️ {title}\n")
    run_t.bold = True
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(0x92, 0x40, 0x0E)

    run_b = p.add_run(text)
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = RGBColor(0x78, 0x35, 0x0F)

    doc.add_paragraph()  # spacing


def generate_phase_5_docx():
    doc = Document()

    # Page Margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # Title Banner
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(2)
    run_t = title_p.add_run("PHASE 5 — MASTER IMPLEMENTATION GUIDE")
    run_t.bold = True
    run_t.font.size = Pt(22)
    run_t.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)  # Indigo-600

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Virtual Gym Buddy — Conversational Fitness Intelligence & Multi-Domain Context Engine")
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    # Callout Notice
    add_callout_box(
        doc,
        "Virtual Gym Buddy provides personalized fitness, squat form, and nutrition coaching grounded in authenticated user context. "
        "It includes strict medical safety boundaries and a seamless deterministic fallback engine when external LLMs are unavailable.",
        title="PHASE 5 SCOPE & SAFETY BOUNDARIES"
    )

    # Sections
    content_sections = [
        ("1. Purpose", [
            "Phase 5 introduces a real conversational Virtual Gym Buddy to the AI Gym & Fitness Assistant platform.",
            "Key Objectives:",
            "• Grounded Fitness Intelligence: Answers fitness queries using actual authenticated user data (Profile, Phase 3 Weekly Workouts, Phase 4 Nutrition).",
            "• User Data Isolation: Guarantees strict multi-tenant isolation where User A can never access User B's profile, workout logs, nutrition targets, or chat history.",
            "• Medical Safety & Boundaries: Detects diagnostic or injury-related prompts and enforces safe boundaries with immediate medical disclaimers.",
            "• Dual-Engine Reliability: Integrates external LLMs (Gemini / OpenAI) when API keys exist, while guaranteeing 100% operational uptime via a Deterministic Expert Fallback Engine."
        ]),
        ("2. Architecture", [
            "End-to-End System Topology:",
            "React / Next.js Frontend (/buddy) ──> FastAPI (/buddy/chat) ──> JWT Authentication ──> BuddyContextEngine ──> BuddyService (LLM / Fallback) ──> PostgreSQL Persistence (buddy_messages).",
            "",
            "Key Architectural Components:",
            "• Next.js Chat Interface: Real-time chat UI with quick prompt chips, provider badges, loading indicators, and active context sidebar.",
            "• BuddyContextEngine: Service that aggregates multi-domain data without granting external LLMs direct SQL database access.",
            "• BuddyService: Core coordinator handling medical safety checks, LLM API dispatch, fallback execution, and PostgreSQL chat persistence.",
            "• BuddyMessage Model: Database table for persistent, user-scoped chat history."
        ]),
        ("3. Data Flow", [
            "Step-by-Step Execution Sequence:",
            "1. Client Request: User sends a message via POST /buddy/chat with Authorization Bearer JWT token.",
            "2. Auth Verification: FastAPI dependency verifies JWT token and extracts authenticated User object.",
            "3. Safety Boundary Check: BuddyService scans message against medical/injury keywords. If triggered, returns immediate safety boundary response.",
            "4. Context Assembly: BuddyContextEngine queries PostgreSQL for User Profile, Phase 3 Weekly Performance Report, and Phase 4 Nutrition Daily Summary.",
            "5. Engine Dispatch: If GEMINI_API_KEY or OPENAI_API_KEY is present, formats system instructions + context block and calls HTTP API. On error or missing key, executes Deterministic Expert Engine.",
            "6. Persistence & Response: Both user message and assistant reply (with provider tag) are persisted to PostgreSQL and returned to the client."
        ]),
        ("4. Context Construction", [
            "Multi-Domain Context Engine (BuddyContextEngine):",
            "• Profile Context: Age, gender, height (cm), weight (kg), fitness goal, activity level, dietary preference.",
            "• Workout & Performance Context: 7-day workout count, total reps, average form score, top exercise, recurring form issues, strongest improvement area, next week focus.",
            "• Nutrition Context: Calorie target, protein/carbs/fat targets, consumed calories/macros today, remaining budget today, meal log count today, active meal plan goal.",
            "• Graceful Missing Data Handling: Users with 0 workouts or 0 nutrition logs receive clean empty context structures without throwing database exceptions."
        ]),
        ("5. LLM Integration & Deterministic Fallback", [
            "Dual-Engine Execution Strategy:",
            "• LLM Provider: Standard urllib HTTP request to Gemini 1.5 Flash or OpenAI GPT-4o-mini with 10s timeout.",
            "• Strict System Prompt: Instructs LLM to rely strictly on provided user context and NEVER fabricate missing workout records or nutrition numbers.",
            "• Deterministic Expert Engine: Rule-based fallback system that analyzes question intent (performance, form, next workout, nutrition, motivation) and returns personalized answers based on gathered context when LLM is unavailable."
        ]),
        ("6. API Contract", [
            "Endpoints Specification:",
            "• POST /buddy/chat",
            "  Request Body: { \"message\": \"...\", \"include_history\": true }",
            "  Response: { \"message\": \"...\", \"provider\": \"llm_gemini|llm_openai|deterministic_buddy_fallback|system_safety\", \"disclaimer\": \"...\", \"context_summary\": {...}, \"timestamp\": \"...\" }",
            "• GET /buddy/history",
            "  Response: { \"user_id\": 1, \"messages\": [ { \"id\": 1, \"sender\": \"user|assistant\", \"content\": \"...\", \"provider\": \"...\", \"created_at\": \"...\" } ] }",
            "• DELETE /buddy/history",
            "  Response: { \"message\": \"Chat history cleared successfully\", \"deleted_count\": 5 }"
        ]),
        ("7. Testing & Verification", [
            "Automated Test Coverage (13 Dedicated Phase 5 Tests + 27 Regression Tests = 40 Total):",
            "• Test 01: Authenticated chat request succeeds (200 OK).",
            "• Test 02: Unauthenticated request rejected (401 Unauthorized).",
            "• Test 03: Cross-user data isolation verified.",
            "• Test 04: Profile context correctly scoped.",
            "• Test 05: Phase 3 workout performance context correctly scoped.",
            "• Test 06: Phase 4 nutrition context correctly scoped.",
            "• Test 07: Missing user data handled safely.",
            "• Test 08: Missing LLM API key falls back to deterministic engine.",
            "• Test 09: LLM API network failure handles gracefully.",
            "• Test 10: Malformed LLM response handles gracefully.",
            "• Test 11: Medical/injury prompts trigger safety disclaimers.",
            "• Test 12: Anti-fabrication verification (no fake statistics created).",
            "• Test 13: Chat history persistence and deletion verified.",
            "• Frontend Verification: Next.js production build (`npm run build`) compiled cleanly with 0 TypeScript/lint errors."
        ]),
    ]

    for section_title, paragraphs in content_sections:
        h = doc.add_heading(section_title, level=1)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)

        for line in paragraphs:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15

            if line.startswith("• "):
                p.paragraph_format.left_indent = Inches(0.25)
                run_bullet = p.add_run("• ")
                run_bullet.bold = True
                run_bullet.font.color.rgb = RGBColor(0x4F, 0x46, 0xE5)
                run_text = p.add_run(line[2:])
            elif line.startswith("Key ") or line.endswith(":") or line.startswith("End-to-End"):
                run_text = p.add_run(line)
                run_text.bold = True
                run_text.font.color.rgb = RGBColor(0x11, 0x18, 0x27)
            else:
                run_text = p.add_run(line)

    output_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "PHASE_5_MASTER_IMPLEMENTATION_GUIDE.docx")
    doc.save(output_path)
    print(f"Document successfully created at: {output_path}")


if __name__ == "__main__":
    generate_phase_5_docx()
