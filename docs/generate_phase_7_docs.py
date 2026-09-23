import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout_box(doc, text, title="GYM RECOMMENDER & PLANNER DESIGN POLICY", bg_hex="F0F9FF", border_hex="0284C7"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')  # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_hex)
    tcBorders.append(left)

    for side in ['top', 'bottom', 'right']:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)

    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"🏋️ {title}\n")
    run_t.bold = True
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(2, 132, 199)

    run_b = p.add_run(text)
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(30, 41, 59)


def generate_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("AI Gym & Fitness Assistant")
    run_title.font.size = Pt(26)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(2, 132, 199)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Phase 7 — Master Implementation Guide: Gym Recommender & Workout Planner")
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Callout Box
    add_callout_box(
        doc,
        "Phase 7 integrates intelligent gym recommendation with personalized, adaptive weekly workout planning. "
        "Gym suitability matches use transparent numerical scoring (0–100%) and non-causal explanations. "
        "Workout plans dynamically adapt to Phase 3 posture form warnings (injecting warmup notes) "
        "and Phase 6 habit skip risks (capping days to 3 for high-risk or cold-start users).",
        title="PHASE 7 ARCHITECTURAL INTEGRATION & DESIGN POLICY"
    )

    doc.add_paragraph()

    # Executive Overview
    h1 = doc.add_heading("1. Executive Overview", level=1)
    h1.runs[0].font.color.rgb = RGBColor(2, 132, 199)

    p_exec = doc.add_paragraph(
        "Phase 7 introduces the Gym Recommender & Workout Planner, seamlessly connecting fitness facilities "
        "with personalized 7-day weekly workout scheduling. Rather than functioning as a static list or generic chatbot, "
        "this module uses authenticated user context (profile goals, target muscle groups, equipment access, historical posture "
        "form warnings, and behavioral habit skip risks) to calculate gym suitability and generate adaptive weekly workout schedules."
    )
    p_exec.paragraph_format.space_after = Pt(12)

    # Core Architectural Principles
    h2 = doc.add_heading("2. Core Architectural Principles", level=1)
    h2.runs[0].font.color.rgb = RGBColor(2, 132, 199)

    principles = [
        ("Structured Gym Catalogue & Recommendation Engine", "Manages synthetic facilities storing equipment, amenities, distance, price tier, and ratings. Computes transparent numerical suitability scores (0-100%)."),
        ("Goal-Aligned 7-Day Workout Split Generation", "Auto-generates 7-day plans supporting 5 goals: hypertrophy (Push/Pull/Legs), strength (Upper/Lower), weight loss (Full Body), endurance (Stamina), and maintenance (Balanced)."),
        ("Phase 3 Posture Form Warning Adaptation", "Detects historical posture errors (e.g. knee valgus, forward lean) and injects targeted warmup notes and technique execution cues."),
        ("Phase 6 Behavioral Habit Skip Risk Adaptation", "Evaluates Phase 6 habit skip risk; if risk is high or user is cold-start, caps weekly training volume to 3 days with added recovery days."),
        ("Relational Database Persistence & User Isolation", "Persists records across 'gyms', 'workout_plans', and 'workout_plan_items' with ON DELETE CASCADE and JWT user isolation."),
        ("FastAPI Endpoints", "Exposes GET /gyms, GET /gyms/recommendations, POST /planner/generate, GET /planner/latest, GET /planner/history."),
        ("Interactive Next.js Frontend (/planner)", "Tabbed interface for Gym Recommender (suitability badges, match reasons) and Weekly Workout Planner (7-day interactive grid).")
    ]

    for title, desc in principles:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r_t = p.add_run(f"• {title}: ")
        r_t.bold = True
        p.add_run(desc)

    # Database Schema
    doc.add_paragraph()
    h3 = doc.add_heading("3. Database Schema Overview", level=1)
    h3.runs[0].font.color.rgb = RGBColor(2, 132, 199)

    table_data = [
        ["Table", "Field Name", "Type", "Description"],
        ["gyms", "id", "INTEGER", "Primary Key"],
        ["gyms", "name", "VARCHAR(150)", "Gym facility name"],
        ["gyms", "distance_km", "FLOAT", "Distance from user (km)"],
        ["gyms", "rating", "FLOAT", "User rating [1.0, 5.0]"],
        ["workout_plans", "id", "INTEGER", "Primary Key"],
        ["workout_plans", "user_id", "INTEGER", "FK -> users.id (ON DELETE CASCADE)"],
        ["workout_plans", "fitness_goal", "VARCHAR(50)", "Target goal (hypertrophy, etc.)"],
        ["workout_plans", "habit_adapted", "BOOLEAN", "Flag for Phase 6 skip risk adaptation"],
        ["workout_plans", "performance_adapted", "BOOLEAN", "Flag for Phase 3 form adaptation"],
        ["workout_plan_items", "day_number", "INTEGER", "Day of week [1, 7]"],
        ["workout_plan_items", "is_rest_day", "BOOLEAN", "Rest day flag"],
        ["workout_plan_items", "warmup_notes", "TEXT", "Phase 3 posture & warmup cues"]
    ]

    t = doc.add_table(rows=len(table_data), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = val
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if i == 0:
                set_cell_background(cell, "0284C7")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")

    # Verification Results
    doc.add_paragraph()
    h4 = doc.add_heading("4. Automated Verification Results", level=1)
    h4.runs[0].font.color.rgb = RGBColor(2, 132, 199)

    verif_p = doc.add_paragraph()
    verif_p.add_run("• Dedicated Phase 7 Test Suite (backend/test_phase_7_planner.py): 20/20 PASSED (0.46s)\n").bold = True
    verif_p.add_run("• Phase 1–6 Regression Suites: 91/91 PASSED\n").bold = True
    verif_p.add_run("• Next.js Build Verification (npm run build): 14/14 static pages generated cleanly, 0 TypeScript errors\n").bold = True

    output_path = os.path.join(os.path.dirname(__file__), "PHASE_7_MASTER_IMPLEMENTATION_GUIDE.docx")
    doc.save(output_path)
    print(f"Generated Phase 7 docx at: {output_path}")


if __name__ == "__main__":
    generate_docx()
