import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_callout_box(doc, text, title="SMART GYM & IOT ARCHITECTURE POLICY", bg_hex="ECFEFF", border_hex="0891B2"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single')
    left.set(qn('w:sz'), '24')  # 3pt
    left.set(qn('w:space'), '0')
    left.set(qn('w:color'), border_hex)
    tcBorders.append(left)

    for side in ['top', 'bottom', 'right']:
        node = OxmlElement(f'w:{side}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)

    tcPr.append(tcBorders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📡 {title}\n")
    run_t.bold = True
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(8, 145, 178)

    run_b = p.add_run(text)
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(30, 41, 59)


def generate_docx():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("AI Gym & Fitness Assistant")
    run_title.font.size = Pt(26)
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(8, 145, 178)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    run_sub = p_sub.add_run("Phase 8 — Master Implementation Guide: Smart Gym Assistant + IoT")
    run_sub.font.size = Pt(15)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Callout Box
    add_callout_box(
        doc,
        "Phase 8 establishes the project's smart equipment intelligence layer. "
        "Smart equipment devices connect via structured MQTT topics (gym/{gym_id}/device/{device_uid}/telemetry, /state, /command). "
        "When physical MQTT brokers are inactive, the system seamlessly defaults to Simulation Mode ('Sample Demo Device / Simulation Mode') "
        "to ensure reliable execution while maintaining full transparency.",
        title="PHASE 8 IOT & SIMULATION POLICY"
    )

    doc.add_paragraph()

    # Executive Overview
    h1 = doc.add_heading("1. Executive Overview", level=1)
    h1.runs[0].font.color.rgb = RGBColor(8, 145, 178)

    p_exec = doc.add_paragraph(
        "Phase 8 introduces the Smart Gym Assistant + IoT, establishing the project's smart equipment intelligence layer. "
        "Rather than operating as a passive dashboard or raw device monitor, this module connects gym equipment devices "
        "(e.g., smart power racks, connected cable crossover stations, assault air bikes) via MQTT communication and REST APIs, "
        "ingests performance telemetry, tracks set intensity and heart rate, provides safe equipment resistance interfaces, "
        "and delivers grounded Smart Gym Assistant heuristics."
    )
    p_exec.paragraph_format.space_after = Pt(12)

    # Core Architectural Principles
    h2 = doc.add_heading("2. Core Architectural Principles", level=1)
    h2.runs[0].font.color.rgb = RGBColor(8, 145, 178)

    principles = [
        ("Structured Equipment Data Model", "Persists IoTDevice, IoTTelemetry, and IoTCommandLog models in PostgreSQL with ON DELETE CASCADE and JWT user isolation."),
        ("MQTT Communication & Simulation Engine", "Supports topic architecture (gym/{gym_id}/device/{device_uid}/telemetry, /state, /command) with Node-RED JSON schema compatibility and transparent Simulation Mode fallback."),
        ("Smart Gym Assistant Heuristics", "Evaluates set intensity, rep counts, and heart rate to recommend 90-120s recovery rest intervals, progressive overload load increases (+2.5kg to +5.0kg), and fatigue alerts."),
        ("Safe Equipment Resistance Interface", "Validates resistance adjustments strictly within physical bounds [0.0 kg, 300.0 kg] for set_resistance, increase_resistance, decrease_resistance, and emergency_stop."),
        ("FastAPI Endpoints Scoped to Authenticated User", "Exposes /iot/devices, /iot/devices/{id}/telemetry, /iot/devices/{id}/command, /iot/assistant/recommendations, and /iot/simulation/generate."),
        ("Interactive Next.js Frontend (/iot)", "Tabbed interface for Connected Devices & Resistance Controls, Telemetry Stream, and Smart Assistant Intelligence.")
    ]

    for title, desc in principles:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        r_t = p.add_run(f"• {title}: ")
        r_t.bold = True
        p.add_run(desc)

    # Database Schema
    doc.add_paragraph()
    h3 = doc.add_heading("3. Database Schema Overview", level=1)
    h3.runs[0].font.color.rgb = RGBColor(8, 145, 178)

    table_data = [
        ["Table", "Field Name", "Type", "Description"],
        ["iot_devices", "id", "INTEGER", "Primary Key"],
        ["iot_devices", "user_id", "INTEGER", "FK -> users.id (ON DELETE CASCADE)"],
        ["iot_devices", "device_uid", "VARCHAR(100)", "Unique equipment UID"],
        ["iot_devices", "is_simulated", "BOOLEAN", "Simulation mode flag (Default True)"],
        ["iot_devices", "current_resistance_kg", "FLOAT", "Current equipment load"],
        ["iot_telemetry", "id", "INTEGER", "Primary Key"],
        ["iot_telemetry", "resistance_kg", "FLOAT", "Telemetry set load [0, 500]"],
        ["iot_telemetry", "repetition_count", "INTEGER", "Reps completed [0, 200]"],
        ["iot_telemetry", "intensity_score", "FLOAT", "Calculated set intensity [0, 100]"],
        ["iot_command_logs", "command_type", "VARCHAR(50)", "Command (set_resistance, etc.)"],
        ["iot_command_logs", "status", "VARCHAR(30)", "Status (simulated_success, sent)"]
    ]

    t = doc.add_table(rows=len(table_data), cols=4)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table_data):
        for j, val in enumerate(row):
            cell = t.cell(i, j)
            cell.text = val
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            if i == 0:
                set_cell_background(cell, "0891B2")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")

    # Verification Results
    doc.add_paragraph()
    h4 = doc.add_heading("4. Automated Verification Results", level=1)
    h4.runs[0].font.color.rgb = RGBColor(8, 145, 178)

    verif_p = doc.add_paragraph()
    verif_p.add_run("• Dedicated Phase 8 Test Suite (backend/test_phase_8_iot.py): 20/20 PASSED (0.757s)\n").bold = True
    verif_p.add_run("• Phase 1–7 Regression Suites: 91/91 PASSED\n").bold = True
    verif_p.add_run("• Total Test Suite Passing: 111/111 PASSED\n").bold = True
    verif_p.add_run("• Next.js Build Verification (npm run build): 15/15 static pages generated cleanly, 0 TypeScript errors\n").bold = True

    output_path = os.path.join(os.path.dirname(__file__), "PHASE_8_MASTER_IMPLEMENTATION_GUIDE.docx")
    doc.save(output_path)
    print(f"Generated Phase 8 docx at: {output_path}")


if __name__ == "__main__":
    generate_docx()
